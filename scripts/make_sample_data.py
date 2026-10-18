"""Generate the bundled sample resume (PDF, DOCX, TXT) and sample job description.

The sample is *illustrative*: PromptShield and ModelRouter are real projects by the author,
the remaining entries (third project, internship bullets) are placeholders so the demo has
something to ask about. Replace it with your own resume in the app.

    python scripts/make_sample_data.py
"""

from __future__ import annotations

from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "src" / "interviewcoach" / "sample_data"

NAME = "Aryaveer Mankotia"
CONTACT = "aryaveer@example.com | github.com/Aryaveer009 | India"

SECTIONS: list[tuple[str, list[tuple[str, list[str]]]]] = [
    (
        "Skills",
        [
            ("", [
                "Languages: Python, SQL, Bash, JavaScript",
                "AI / ML: LLMs, RAG, Prompt Engineering, Embeddings, LangChain, Hugging Face",
                "Security: Cybersecurity, OWASP, Prompt Injection, Threat Modeling, Burp Suite, Nmap",
                "Tools: Git, Docker, Linux, pytest, FastAPI",
            ])
        ],
    ),
    (
        "Projects",
        [
            ("PromptShield | Python, Typer, GitHub Actions", [
                "Built an open-source CLI that runs 48 attacks across 7 categories against an LLM "
                "chatbot or RAG app and reports pass/fail mapped to the OWASP Top 10 for LLM "
                "Applications.",
                "Designed canary-secret detection that catches system-prompt leaks even when the "
                "secret is reversed, spaced out or base64-encoded.",
                "Added a mandatory ownership flag so the tool refuses to run against targets the "
                "user has not explicitly claimed.",
                "Set up GitHub Actions CI for linting, type-checking and tests.",
            ]),
            ("ModelRouter | Python, FastAPI, Docker", [
                "Built an OpenAI-compatible proxy that routes easy queries to a cheap model and "
                "hard queries to a strong model to cut LLM costs.",
                "Implemented a query-difficulty classifier and per-request cost tracking behind a "
                "REST API.",
                "Containerized the service with Docker and added a Streamlit dashboard for spend "
                "and routing decisions.",
            ]),
            ("Campus Notes Q&A (sample project) | Python, RAG, SQLite", [
                "Built a retrieval-augmented generation assistant that answers questions over "
                "lecture notes using embeddings and a FAISS vector index.",
                "Chunked documents and stored metadata in SQLite so answers could cite their "
                "source pages.",
            ]),
        ],
    ),
    (
        "Experience",
        [
            ("Cybersecurity Intern | Security Operations (sample entry)", [
                "Triaged security alerts and documented findings for the team.",
                "Ran vulnerability scans with Nmap and summarized remediation steps for engineers.",
            ]),
        ],
    ),
    ("Education", [("", ["B.Tech, Computer Science"])]),
]

JOB_DESCRIPTION = """Junior AI Security Engineer
Acme Labs (fictional company, synthetic job description)

About the role
You will help us ship LLM-powered features safely. You will build tooling that tests our
chatbots and retrieval pipelines for security weaknesses and work with engineers to fix them.

Responsibilities
- Build and maintain automated security tests for LLM and RAG applications
- Write Python services and internal tools exposed through REST APIs
- Analyse data in SQL to find patterns in failed attacks and incidents
- Document threats and help the team with threat modeling for new features

Requirements
- Strong Python
- Solid SQL
- Hands-on experience with LLMs and retrieval-augmented generation (RAG)
- Understanding of OWASP and prompt injection attacks
- Experience building REST APIs
- Familiarity with Docker and Kubernetes
- Git and CI/CD workflows

Nice to have
- PostgreSQL
- AWS
- Terraform
- Penetration testing experience
"""


def _txt() -> str:
    lines = [NAME, CONTACT, ""]
    for heading, entries in SECTIONS:
        lines.append(heading.upper())
        for title, bullets in entries:
            if title:
                lines.append(title)
            for b in bullets:
                lines.append(f"- {b}" if title else b)
        lines.append("")
    return "\n".join(lines)


def _docx(path: Path) -> None:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.add_heading(NAME, level=0)
    doc.add_paragraph(CONTACT)
    for heading, entries in SECTIONS:
        doc.add_heading(heading, level=1)
        for title, bullets in entries:
            if title:
                p = doc.add_paragraph()
                run = p.add_run(title)
                run.bold = True
                run.font.size = Pt(11)
            for b in bullets:
                if title:
                    doc.add_paragraph(b, style="List Bullet")
                else:
                    doc.add_paragraph(b)
    doc.save(path)


def _pdf(path: Path) -> None:
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Title"], fontSize=20, alignment=0, spaceAfter=2)
    meta = ParagraphStyle("meta", parent=styles["Normal"], fontSize=9, textColor="#444444")
    sec = ParagraphStyle("sec", parent=styles["Heading2"], fontSize=11.5, spaceBefore=10,
                         spaceAfter=3, textColor="#1f3a8a")
    entry = ParagraphStyle("entry", parent=styles["Normal"], fontName="Helvetica-Bold",
                           fontSize=10, spaceBefore=4)
    body = ParagraphStyle("body", parent=styles["Normal"], fontSize=9.5, leading=12.5)
    bullet = ParagraphStyle("bullet", parent=body, leftIndent=12, bulletIndent=2)

    story = [Paragraph(escape(NAME), h1), Paragraph(escape(CONTACT), meta)]
    for heading, entries in SECTIONS:
        story.append(Paragraph(heading.upper(), sec))
        for title, bullets in entries:
            if title:
                story.append(Paragraph(escape(title), entry))
            for b in bullets:
                if title:
                    story.append(Paragraph(escape(b), bullet, bulletText="•"))
                else:
                    story.append(Paragraph(escape(b), body))
        story.append(Spacer(1, 2 * mm))
    SimpleDocTemplate(
        str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=15 * mm,
        bottomMargin=15 * mm, title=f"{NAME} - sample resume", author=NAME,
    ).build(story)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "resume.txt").write_text(_txt(), encoding="utf-8")
    (OUT / "job_description.txt").write_text(JOB_DESCRIPTION, encoding="utf-8")
    _docx(OUT / "resume.docx")
    _pdf(OUT / "resume.pdf")
    print(f"wrote sample files to {OUT}")


if __name__ == "__main__":
    main()
