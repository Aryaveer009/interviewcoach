"""Resume and job-description parsing.

Parsing is rule-based on purpose: it is deterministic, works offline, and the
resume never has to leave the machine just to be split into sections.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from .models import Experience, JobDescription, Project, ResumeProfile
from .skills import extract_skills

BULLET_CHARS = "•▪◦·●‣⁃-–—*"
_BULLET_RE = re.compile(rf"^\s*[{re.escape(BULLET_CHARS)}]\s+")

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "summary": ("summary", "profile", "objective", "about me"),
    "experience": ("experience", "work experience", "professional experience", "internships",
                   "internship experience", "employment"),
    "projects": ("projects", "personal projects", "academic projects", "selected projects"),
    "education": ("education", "academics", "academic background"),
    "skills": ("skills", "technical skills", "core skills", "skills & tools", "tech stack"),
    "other": ("certifications", "achievements", "awards", "publications", "coursework",
              "extracurriculars", "activities", "leadership"),
}
_HEADING_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
LINK_RE = re.compile(
    r"(?:https?://)?(?:www\.)?(?:github\.com|linkedin\.com/in|gitlab\.com)/[\w\-./]+", re.I
)


class ParseError(ValueError):
    """Raised when a file cannot be turned into text."""


# --------------------------------------------------------------------------- file -> text


def extract_text(data: bytes, filename: str) -> str:
    """Return plain text from PDF / DOCX / TXT / MD bytes."""
    suffix = Path(filename).suffix.lower()
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        elif suffix == ".docx":
            from docx import Document

            doc = Document(io.BytesIO(data))
            # python-docx drops list glyphs, so re-add a bullet for list-styled paragraphs
            parts = [
                f"• {p.text}" if (p.style is not None and (p.style.name or "").startswith("List"))
                else p.text
                for p in doc.paragraphs
            ]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif suffix in (".txt", ".md", ""):
            text = data.decode("utf-8", errors="replace")
        else:
            raise ParseError(f"Unsupported file type {suffix!r}; use PDF, DOCX or TXT.")
    except ParseError:
        raise
    except Exception as exc:  # corrupt / encrypted files
        raise ParseError(f"Could not read {filename}: {exc}") from exc
    text = text.replace("\r\n", "\n").replace("\x00", "")
    # PDF bullet glyphs often extract as DEL or private-use characters
    text = re.sub(r"[\x7f\uf000-\uf0ff]", "•", text)
    if not text.strip():
        raise ParseError(
            f"No text found in {filename}. Scanned/image-only PDFs are not supported (no OCR)."
        )
    return text


# --------------------------------------------------------------------------- sections


def _heading_key(line: str) -> str | None:
    cleaned = re.sub(r"[:\-–—_=*#|]+", " ", line).strip().lower()
    cleaned = re.sub(r"\s+", " ", cleaned)
    if not cleaned or len(cleaned) > 40:
        return None
    return _HEADING_LOOKUP.get(cleaned)


def split_sections(text: str) -> dict[str, list[str]]:
    """Group lines under their section heading; everything before the first is ``header``."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        key = None if _BULLET_RE.match(line) else _heading_key(line)
        if key:
            current = key
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(line.strip())
    return sections


def _is_bullet(line: str) -> bool:
    return bool(_BULLET_RE.match(line))


def _strip_bullet(line: str) -> str:
    return _BULLET_RE.sub("", line, count=1).strip()


def _entry_name(title: str) -> str:
    return re.split(r"\s+[|–—]\s+|\s+-\s+|\s*\(", title, maxsplit=1)[0].strip(" :,")


def _looks_like_title(line: str, prev_bullet: str | None) -> bool:
    if prev_bullet is not None and not prev_bullet.rstrip().endswith((".", "!", "?", ";")):
        # previous bullet has no terminal punctuation: this is a wrapped continuation
        return False
    if line[:1].islower():
        return False
    return len(line) <= 110 and not line.rstrip().endswith(".")


def _parse_entries(lines: list[str]) -> list[tuple[str, list[str], list[str]]]:
    """Return ``(title, description_lines, bullets)`` per entry in a section."""
    entries: list[tuple[str, list[str], list[str]]] = []
    for line in lines:
        cur = entries[-1] if entries else None
        if _is_bullet(line):
            if cur is None:
                entries.append(("(untitled)", [], []))
                cur = entries[-1]
            cur[2].append(_strip_bullet(line))
            continue
        prev_bullet = cur[2][-1] if cur and cur[2] else None
        if cur is not None and cur[2] and not _looks_like_title(line, prev_bullet):
            cur[2][-1] = f"{cur[2][-1]} {line}".strip()  # wrapped bullet continuation
        elif cur is None or cur[2]:
            entries.append((line, [], []))
        else:
            cur[1].append(line)  # description line under a title, before bullets
    return entries


def _parse_projects(lines: list[str]) -> list[Project]:
    projects = []
    for title, desc, bullets in _parse_entries(lines):
        body = " ".join([title, *desc, *bullets])
        projects.append(
            Project(
                name=_entry_name(title) or title,
                description=" ".join(desc),
                bullets=bullets,
                skills=extract_skills(body),
            )
        )
    return projects


def _parse_experience(lines: list[str]) -> list[Experience]:
    out = []
    for title, desc, bullets in _parse_entries(lines):
        body = " ".join([title, *desc, *bullets])
        out.append(
            Experience(title=title, bullets=[*desc, *bullets], skills=extract_skills(body))
        )
    return out


# --------------------------------------------------------------------------- public API


def parse_resume(text: str) -> ResumeProfile:
    sections = split_sections(text)
    header = sections.get("header", [])
    name = None
    for line in header[:3]:
        if len(line) <= 50 and "@" not in line and not re.search(r"\d{4}", line):
            name = line.strip()
            break
    email_match = EMAIL_RE.search(text)
    links = sorted({m.group(0).rstrip(".,)") for m in LINK_RE.finditer(text)})
    skills_text = "\n".join(sections.get("skills", []))
    return ResumeProfile(
        name=name,
        email=email_match.group(0) if email_match else None,
        links=links,
        skills=extract_skills(text),
        skills_listed=extract_skills(skills_text),
        projects=_parse_projects(sections.get("projects", [])),
        experience=_parse_experience(sections.get("experience", [])),
        education=sections.get("education", []),
        raw_text=text,
    )


_REQUIRED_HEADS = ("requirement", "qualification", "must have", "must-have", "what you'll need",
                   "what you will need", "what we're looking for", "you have", "minimum")
_PREFERRED_HEADS = ("nice to have", "nice-to-have", "preferred", "bonus", "a plus", "good to have",
                    "desirable")
_RESP_HEADS = ("responsibilit", "what you'll do", "what you will do", "the role", "your role",
               "day to day", "you will")


def _jd_section_kind(line: str) -> str | None:
    low = line.strip().lower().rstrip(":")
    if len(low) > 60 or _is_bullet(line):
        return None
    if any(h in low for h in _PREFERRED_HEADS):
        return "preferred"
    if any(h in low for h in _REQUIRED_HEADS):
        return "required"
    if any(h in low for h in _RESP_HEADS):
        return "responsibilities"
    return None


def parse_job_description(text: str) -> JobDescription:
    """Split a JD into required / preferred skills and responsibilities.

    Skills mentioned outside a "preferred" section are treated as required, and a
    skill that appears in both is kept as required.
    """
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    kind = "intro"
    buckets: dict[str, list[str]] = {"intro": [], "required": [], "preferred": [],
                                      "responsibilities": []}
    for line in lines:
        new_kind = _jd_section_kind(line)
        if new_kind:
            kind = new_kind
            continue
        buckets[kind].append(line)
    title = re.sub(r"^(job title|title|role)\s*[:\-]\s*", "", lines[0], flags=re.I) if lines else None
    required = extract_skills("\n".join(buckets["intro"] + buckets["required"]
                                        + buckets["responsibilities"]))
    preferred = [s for s in extract_skills("\n".join(buckets["preferred"])) if s not in required]
    return JobDescription(
        title=title,
        required_skills=required,
        preferred_skills=preferred,
        responsibilities=[_strip_bullet(ln) for ln in buckets["responsibilities"]],
        raw_text=text,
    )
